import csv
import json
from pathlib import Path
from typing import Dict, List, Optional, Callable
from threading import Lock
from openpyxl import load_workbook, Workbook
from src.utils.logger import get_logger
from src.utils.cache import LocalCache
from src.utils.validators import LinkedInValidator, DataValidator
from src.services.search_service import SearchService
from config.settings import DATA_DIR

logger = get_logger(__name__)

class DataProcessor:
    def __init__(self):
        self.cache = LocalCache(DATA_DIR)
        self.search_service = SearchService()
        self.lock = Lock()
        self.is_running = False

    def load_spreadsheet(self, file_path: str) -> Optional[List[Dict]]:
        try:
            file_path = Path(file_path)

            if file_path.suffix.lower() == '.csv':
                return self._load_csv(file_path)
            elif file_path.suffix.lower() in ['.xlsx', '.xls']:
                return self._load_excel(file_path)
            else:
                logger.error(f"Formato não suportado: {file_path.suffix}")
                return None

        except Exception as e:
            logger.error(f"Erro ao carregar planilha: {e}")
            return None

    def _load_csv(self, file_path: Path) -> Optional[List[Dict]]:
        try:
            data = []
            with open(file_path, 'r', encoding='utf-8') as f:
                reader = csv.DictReader(f)
                for row in reader:
                    normalized = {k.lower().strip(): v for k, v in row.items()}
                    data.append(normalized)
            logger.info(f"CSV carregado: {len(data)} linhas")
            return data
        except Exception as e:
            logger.error(f"Erro ao carregar CSV: {e}")
            return None

    def _load_excel(self, file_path: Path) -> Optional[List[Dict]]:
        try:
            wb = load_workbook(file_path)
            ws = wb.active
            data = []
            headers = [cell.value for cell in ws[1]]
            headers = [str(h).lower().strip() if h else "" for h in headers]

            for row in ws.iter_rows(min_row=2, values_only=True):
                row_dict = {headers[i]: row[i] for i in range(len(headers)) if i < len(row)}
                data.append(row_dict)

            logger.info(f"Excel carregado: {len(data)} linhas")
            return data
        except Exception as e:
            logger.error(f"Erro ao carregar Excel: {e}")
            return None

    def enrich_contacts(
        self,
        data: List[Dict],
        progress_callback: Optional[Callable] = None,
        status_callback: Optional[Callable] = None,
    ) -> List[Dict]:
        try:
            self.is_running = True
            output_columns = [
                'linkedin_encontrado',
                'empresa_atual',
                'cargo_atual',
                'score_confianca',
                'status_busca',
                'observacoes',
            ]

            enriched_data = []
            total = len(data)

            for idx, row in enumerate(data):
                if not self.is_running:
                    logger.info("Processamento interrompido pelo usuário")
                    break

                try:
                    enriched_row = {**row}
                    result = self._process_contact(row)
                    enriched_row.update(result)
                    enriched_data.append(enriched_row)

                    if status_callback:
                        status_callback(f"Processado: {row.get('nome', 'N/A')}")

                except Exception as e:
                    logger.error(f"Erro ao processar contato {idx}: {e}")
                    enriched_row = {**row}
                    enriched_row.update({
                        'status_busca': 'erro',
                        'observacoes': str(e),
                    })
                    enriched_data.append(enriched_row)

                finally:
                    if progress_callback:
                        progress_callback((idx + 1) / total * 100)

            self.is_running = False
            logger.info("Enriquecimento concluído")
            return enriched_data

        except Exception as e:
            logger.error(f"Erro no enriquecimento: {e}")
            self.is_running = False
            return data

    def _process_contact(self, row: Dict) -> Dict:
        result = {
            'linkedin_encontrado': '',
            'empresa_atual': '',
            'cargo_atual': '',
            'score_confianca': 0.0,
            'status_busca': 'pendente',
            'observacoes': '',
        }

        is_valid, error_msg = DataValidator.validate_row(row)
        if not is_valid:
            result['status_busca'] = 'erro'
            result['observacoes'] = error_msg
            return result

        nome = DataValidator.clean_text(row.get('nome', ''))
        empresa = DataValidator.clean_text(row.get('empresa', ''))
        cargo = DataValidator.clean_text(row.get('cargo', ''))
        linkedin_url = DataValidator.clean_text(row.get('linkedin', ''))

        if linkedin_url and LinkedInValidator.is_valid_linkedin_url(linkedin_url):
            return self._process_existing_linkedin(linkedin_url, nome, result)

        return self._search_linkedin_profile(nome, empresa, cargo, result)

    def _process_existing_linkedin(
        self,
        url: str,
        name: str,
        result: Dict,
    ) -> Dict:
        try:
            url = LinkedInValidator.normalize_linkedin_url(url)
            cached = self.cache.get(name)

            if cached and cached.get('url') == url:
                return {**result, **cached}

            profile = self.search_service.extract_profile_info(url)

            if profile:
                self.cache.set(name, {
                    'linkedin_encontrado': url,
                    'empresa_atual': profile.get('company', ''),
                    'cargo_atual': profile.get('title', ''),
                    'status_busca': 'encontrado',
                })

                return {
                    **result,
                    'linkedin_encontrado': url,
                    'empresa_atual': profile.get('company', ''),
                    'cargo_atual': profile.get('title', ''),
                    'score_confianca': 1.0,
                    'status_busca': 'encontrado',
                }
            else:
                return {
                    **result,
                    'linkedin_encontrado': url,
                    'status_busca': 'url_invalida',
                    'observacoes': 'URL não retornou informações válidas',
                }

        except Exception as e:
            logger.error(f"Erro ao processar LinkedIn existente: {e}")
            return {
                **result,
                'status_busca': 'erro',
                'observacoes': str(e),
            }

    def _search_linkedin_profile(
        self,
        nome: str,
        empresa: str,
        cargo: str,
        result: Dict,
    ) -> Dict:
        try:
            cached = self.cache.get(nome, empresa, cargo)
            if cached:
                logger.debug(f"Usando cache para {nome}")
                return {**result, **cached}

            search_results = self.search_service.search_linkedin_profile(nome, empresa, cargo)

            if not search_results:
                return {
                    **result,
                    'status_busca': 'nao_encontrado',
                    'observacoes': 'Nenhum perfil encontrado',
                }

            best_match = None
            best_score = 0.0

            for search_result in search_results:
                profile = self.search_service.extract_profile_info(search_result['url'])
                if not profile:
                    continue

                is_match, score = self.search_service.validate_profile_match(
                    profile,
                    nome,
                    empresa,
                    cargo,
                )

                if score > best_score:
                    best_score = score
                    best_match = (profile, score)

            if best_match:
                profile, score = best_match
                cached_data = {
                    'linkedin_encontrado': profile['url'],
                    'empresa_atual': profile.get('company', ''),
                    'cargo_atual': profile.get('title', ''),
                    'score_confianca': round(score, 3),
                    'status_busca': 'encontrado' if score >= 0.5 else 'duvidoso',
                }

                self.cache.set(nome, cached_data, empresa, cargo)
                return {**result, **cached_data}

            return {
                **result,
                'status_busca': 'nao_encontrado',
                'observacoes': 'Nenhum perfil com score suficiente',
            }

        except Exception as e:
            logger.error(f"Erro na busca de perfil para {nome}: {e}")
            return {
                **result,
                'status_busca': 'erro',
                'observacoes': str(e),
            }

    def stop_processing(self) -> None:
        self.is_running = False
        logger.info("Parada de processamento solicitada")

    def export_spreadsheet(self, data: List[Dict], output_path: str) -> bool:
        try:
            output_path = Path(output_path)
            output_path.parent.mkdir(parents=True, exist_ok=True)

            if output_path.suffix.lower() == '.csv':
                return self._export_csv(data, output_path)
            elif output_path.suffix.lower() in ['.xlsx', '.xls']:
                return self._export_excel(data, output_path)
            else:
                logger.error(f"Formato não suportado: {output_path.suffix}")
                return False

        except Exception as e:
            logger.error(f"Erro ao exportar planilha: {e}")
            return False

    def _export_csv(self, data: List[Dict], output_path: Path) -> bool:
        try:
            if not data:
                logger.warning("Nenhum dado para exportar")
                return False

            fieldnames = list(data[0].keys()) if data else []

            with open(output_path, 'w', newline='', encoding='utf-8') as f:
                writer = csv.DictWriter(f, fieldnames=fieldnames)
                writer.writeheader()
                writer.writerows(data)

            logger.info(f"CSV exportado: {output_path}")
            return True
        except Exception as e:
            logger.error(f"Erro ao exportar CSV: {e}")
            return False

    def _export_excel(self, data: List[Dict], output_path: Path) -> bool:
        try:
            if not data:
                logger.warning("Nenhum dado para exportar")
                return False

            wb = Workbook()
            ws = wb.active
            fieldnames = list(data[0].keys())

            ws.append(fieldnames)

            for row_data in data:
                row = [row_data.get(field, "") for field in fieldnames]
                ws.append(row)

            wb.save(output_path)
            logger.info(f"Excel exportado: {output_path}")
            return True
        except Exception as e:
            logger.error(f"Erro ao exportar Excel: {e}")
            return False
